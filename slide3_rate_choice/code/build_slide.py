import pickle, sys, numpy as np, pandas as pd, matplotlib
matplotlib.use('Agg'); import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = sys.argv[1]
D = pickle.load(open('rate_scen.pkl', 'rb')); RES, PR = D['RES'], D['PR']
for _d in (RES, D['RES1']):
    for _v in _d.values(): _v.attrs = {}
COL = {15.0: '#C8661A', 16.0: '#008F99', 17.0: '#5B3F8C'}
LAB = {15.0: '15 %', 16.0: '16 % (база)', 17.0: '17 %'}
LIME = RGBColor(0xBF, 0xD7, 0x30); INK = RGBColor(0x1A, 0x1A, 0x1A); GREEN_TXT = RGBColor(0x6E, 0x9A, 0x1E)
f = lambda x, d=2: f'{x:.{d}f}'.replace('.', ',').replace('-', '−')

# ---------- 2x2 міні-графіки
plt.rcParams.update({'font.family': 'Liberation Sans', 'font.size': 9, 'axes.spines.top': False, 'axes.spines.right': False,
                     'axes.edgecolor': '#9A9A9A', 'axes.grid': True, 'grid.color': '#E3E3E3', 'grid.linewidth': .6,
                     'xtick.color': '#555555', 'ytick.color': '#555555', 'axes.titleweight': 'normal'})
q = list(RES[16.0].index); xi = np.arange(len(q)); xl = [s.replace('20', '’', 1) if False else s for s in q]
panels = [('r', 'Облікова ставка, % (середня за квартал)', []),
          ('cpi_yoy', 'Інфляція ІСЦ, % р/р', ['2027Q4', '2028Q3']),
          ('gdp_yoy', 'Реальний ВВП, % р/р', ['2027Q4']),
          ('eur_uah', 'Курс гривні до євро, грн/EUR', ['2028Q3'])]
fig, axs = plt.subplots(2, 2, figsize=(6.45, 4.55), dpi=300)
for ax, (v, t, marks) in zip(axs.ravel(), panels):
    for m in marks: ax.axvline(q.index(m), color='#BFD730', lw=6, alpha=.35, zorder=0)
    for k in (15.0, 17.0, 16.0):
        ax.plot(xi, RES[k][v], color=COL[k], lw=2.4 if k == 16.0 else 1.8, label=LAB[k], zorder=3 if k == 16.0 else 2)
        for m in marks:
            j = q.index(m); ax.plot(j, RES[k][v].iloc[j], 'o', ms=4.5, color=COL[k], mec='white', mew=.8, zorder=4)
    ax.set_title(t, fontsize=9.5, loc='left', color='#1A1A1A', pad=4)
    ax.set_xticks(xi, [s[2:4] + s[4:] if False else s for s in q], fontsize=6.6, rotation=0)
    ax.set_xticklabels([s.replace('Q', '\nQ') for s in q])
    ax.tick_params(length=0, labelsize=7)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda y, _: f(y, 1)))
h, l = axs[0, 0].get_legend_handles_labels(); order = [l.index(LAB[k]) for k in (15.0, 16.0, 17.0)]
fig.legend([h[i] for i in order], [f'рішення 29.10: {l[i]}' for i in order], loc='upper center', ncol=3, frameon=False, fontsize=8.5, bbox_to_anchor=(.5, 1.0))
fig.tight_layout(rect=(0, 0, 1, .95), h_pad=1.2, w_pad=1.6)
fig.savefig(f'{OUT}/mini_charts.png', dpi=300); plt.close(fig)

# ---------- PPTX: один слайд 16:9
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
def tb(x, y, w, h, text, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT, italic=False, anchor=MSO_ANCHOR.TOP):
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = b.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        for seg, sb in (line if isinstance(line, list) else [(line, bold)]):
            r = p.add_run(); r.text = seg; r.font.size = Pt(size); r.font.bold = sb; r.font.italic = italic
            r.font.color.rgb = color; r.font.name = 'Arial'
    return b
tb(0.42, 0.18, 12.4, 0.75, 'Чому 16 %, а не 15 % чи 17 %', size=30)
ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.42), Inches(0.92), Inches(12.45), Inches(0.07))
ln.fill.solid(); ln.fill.fore_color.rgb = LIME; ln.line.fill.background(); ln.shadow.inherit = False
tb(0.42, 1.05, 6.6, 0.32, 'Три прогони FINEX-QPM на 8 кварталів', size=15)
tb(0.42, 1.36, 6.6, 0.3, 'рішення 29.10.2026 (чинність з 30.10): 15 %, 16 % або 17 %; далі ставка — за правилом (57)', size=10, color=GREEN_TXT)
s.shapes.add_picture(f'{OUT}/mini_charts.png', Inches(0.42), Inches(1.68), Inches(6.45), Inches(4.55))
tb(5.25, 6.12, 1.7, 0.2, 'Джерело: модель команди', size=8, italic=True, align=PP_ALIGN.RIGHT)

# таблиця показників
rows = [('Середня ставка 2026Q4, %', 'r', '2026Q4', 2),
        ('Реальна ставка 2026Q4 (ex ante), %', 'rr', '2026Q4', 1),
        ('Інфляція ІСЦ 2027Q4, % р/р', 'cpi_yoy', '2027Q4', 2),
        ('Інфляція ІСЦ 2028Q3, % р/р', 'cpi_yoy', '2028Q3', 2),
        ('ВВП 2027Q4, % р/р', 'gdp_yoy', '2027Q4', 2),
        ('Курс 2028Q3, грн/EUR', 'eur_uah', '2028Q3', 2),
        ('Резерви 2028Q3, міс. імпорту', 'MI', '2028Q3', 2),
        ('Імовірність утримати резерви ≥ 3 міс. імпорту в усі 8 кв., %', None, None, 1)]
def cell(c, text, size=11, bold=False, fill=None, align=PP_ALIGN.CENTER, color=INK):
    c.text = ''; p = c.text_frame.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = 'Arial'; r.font.color.rgb = color
    c.vertical_anchor = MSO_ANCHOR.MIDDLE; c.margin_left = c.margin_right = Inches(.06); c.margin_top = c.margin_bottom = Inches(.02)
    c.fill.solid(); c.fill.fore_color.rgb = fill or RGBColor(0xFF, 0xFF, 0xFF)
X0, W0 = 7.2, 5.67
T = s.shapes.add_table(len(rows) + 1, 4, Inches(X0), Inches(1.08), Inches(W0), Inches(0.29 * (len(rows) + 1))).table
for j, w in enumerate([2.97, 0.9, 0.9, 0.9]): T.columns[j].width = Inches(w)
cell(T.cell(0, 0), 'Показник', bold=True, fill=LIME, align=PP_ALIGN.LEFT)
for j, k in enumerate((15.0, 16.0, 17.0), 1): cell(T.cell(0, j), f'{int(k)} %', bold=True, fill=LIME)
LIGHT = RGBColor(0xF1, 0xF6, 0xDC)
for i, (name, v, qq, d) in enumerate(rows, 1):
    cell(T.cell(i, 0), name, size=10, align=PP_ALIGN.LEFT)
    for j, k in enumerate((15.0, 16.0, 17.0), 1):
        val = 100 * PR[k]['MI>=3 all8'] if v is None else RES[k].loc[qq, v]
        cell(T.cell(i, j), f(val, d), size=10.5, bold=(k == 16.0), fill=LIGHT if k == 16.0 else None)
for i in range(len(rows) + 1): T.rows[i].height = Inches(0.29)
T.rows[len(rows)].height = Inches(0.42)

# висновки
y2 = 1.08 + 0.29 * len(rows) + 0.42 + 0.17
tb(X0, y2 - 0.02, W0, 0.3, 'Висновок: компроміс між дезінфляцією, ВВП і курсом', size=13, bold=True)
c15, c16, c17 = RES[15.0], RES[16.0], RES[17.0]
dpi = c15.loc['2027Q4', 'cpi_yoy'] - c16.loc['2027Q4', 'cpi_yoy']; dgdp = c15.loc['2027Q4', 'gdp_yoy'] - c16.loc['2027Q4', 'gdp_yoy']
deur = c15.loc['2028Q3', 'eur_uah'] - c16.loc['2028Q3', 'eur_uah']
V = [(15.0, f'Сповільнює дезінфляцію (ІСЦ +{f(dpi,1)} п.п. у 2027Q4) за мізерного виграшу ВВП (+{f(dgdp)} п.п.); слабша гривня (+{f(deur)} грн/EUR) і менші резерви. '
            f'Знижувати при ІСЦ 10 % і очікуваннях 12 % — ризик для довіри; у 2028 р. правило вимагатиме вищої ставки ({f(c15.loc["2028Q3","r"],1)} %).'),
     (16.0, f'Збігається з правилом (57): {f(c16.loc["2026Q4","r"])} % ≈ 16 %. Реальна ставка {f(c16.loc["2026Q4","rr"],1)} % проти нейтральної {f(c16.loc["2026Q4","rr_bar"],1)} % '
            f'достатня, щоб ІСЦ знизився з 10,0 % до {f(c16.loc["2028Q3","cpi_yoy"],1)} % без додаткових втрат ВВП.'),
     (17.0, f'Інфляція нижча лише на {f(-(c17.loc["2027Q4","cpi_yoy"]-c16.loc["2027Q4","cpi_yoy"]),1)} п.п., ціль 5 % на горизонті однаково не досягається; ВВП −{f(-(c17.loc["2027Q4","gdp_yoy"]-c16.loc["2027Q4","gdp_yoy"]))} п.п., '
            f'реальна ставка {f(c17.loc["2026Q4","rr"],1)} %. Ефект короткий: до 2027Q4 ставки всіх прогонів сходяться ({f(c15.loc["2027Q4","r"],1)}–{f(c17.loc["2027Q4","r"],1)} %).')]
VT = s.shapes.add_table(3, 2, Inches(X0), Inches(y2 + 0.3), Inches(W0), Inches(2.0)).table
VT.columns[0].width = Inches(0.62); VT.columns[1].width = Inches(W0 - 0.62)
for i, (k, txt) in enumerate(V):
    hl = k == 16.0
    cell(VT.cell(i, 0), f'{int(k)} %', size=12, bold=True, fill=LIME if hl else RGBColor(0xF2, 0xF2, 0xF2))
    cell(VT.cell(i, 1), txt, size=9, align=PP_ALIGN.LEFT, fill=LIGHT if hl else None)
    VT.rows[i].height = Inches(0.62)

# методична примітка (стиль рамки слайду 2)
nb = tb(0.42, 6.32, 6.62, 0.78, [[('16 % ', True), ('— базовий прогноз (правило (57) дає середню 2026Q4 15,92 %). ', False),
        ('15 % і 17 %', True), (' — той самий прогноз + одноразовий шок правила (57) у 2026Q4: нова ставка діє 63 з 92 днів кварталу, тож середня 2026Q4 '
        'змінюється на ∓0,68 п.п.; з 2027Q1 ставка — за правилом. Імовірність — 5000 бутстреп-траєкторій шоків (без 2022 р.), центрованих на кожному прогоні; '
        'віяло резервів широке через історичні шоки офіційного фінансування.', False)]], size=8.5, italic=True)
nb.line.color.rgb = RGBColor(0x6E, 0x9A, 0x1E); nb.line.width = Pt(1.25)
nb.text_frame.margin_left = nb.text_frame.margin_right = Inches(.08); nb.text_frame.margin_top = nb.text_frame.margin_bottom = Inches(.04)
tb(11.2, 7.08, 1.7, 0.3, 'Жовтень 2026', size=9, bold=True, color=GREEN_TXT, align=PP_ALIGN.RIGHT)
prs.save(f'{OUT}/slide3_chomu_16.pptx')

# ---------- дані для XLSX
with pd.ExcelWriter(f'{OUT}/slide3_rate_scenarios.xlsx') as w:
    summ = pd.DataFrame({LAB[k]: [RES[k].loc[qq, v] if v else 100 * PR[k]['MI>=3 all8'] for (n_, v, qq, d) in rows] for k in RES}, index=[r_[0] for r_ in rows])
    summ.loc['Імовірність резервів ≥ 3 міс. імпорту у 2028Q3, %'] = [100 * PR[k]['MI>=3 2028Q3'] for k in RES]
    summ.loc['Імовірність ІСЦ ≤ 7 % у 2027Q4, %'] = [100 * PR[k]['pi4<=log(1.07) 2027Q4'] for k in RES]
    summ.round(3).rename_axis('показник').to_excel(w, sheet_name='summary')
    names = {'r': 'облікова ставка, % (сер. за кв.)', 'rr': 'реальна ставка ex ante, %', 'rr_bar': 'нейтральна реальна ставка, %', 'pie': 'інфл. очікування, %',
             'cpi_yoy': 'ІСЦ, % р/р', 'pi4': 'ІСЦ, 100·ln р/р', 'gdp_yoy': 'ВВП, % р/р', 'dy': 'ВВП, % кв/кв SAAR', 'y_gap': 'розрив ВВП, %',
             'eur_uah': 'курс, грн/EUR', 'ds_eur': 'девальвація до EUR, % річн.', 'gam': 'премія за гривневі активи, в.п.', 'fa_o': 'потоки резидентів, % ВВП',
             'FXI': 'інтервенції НБУ, % ВВП', 'FXR': 'резерви, % ВВП', 'MI': 'резерви, міс. імпорту', 'OFF': 'офіційне фінансування, % ВВП'}
    for k in RES: RES[k].rename(columns=names).round(3).rename_axis('квартал').to_excel(w, sheet_name=f'run_{int(k)}')
    sens = pd.concat({f'{int(k)} % (±1 п.п. до середньої 2026Q4)': D['RES1'][k][['r', 'cpi_yoy', 'gdp_yoy', 'eur_uah', 'MI']].rename(columns=names) for k in D['RES1']}, axis=1)
    sens.round(3).rename_axis('квартал').to_excel(w, sheet_name='sensitivity_pm1_avg')
print('ok')
